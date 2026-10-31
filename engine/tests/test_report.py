from io import BytesIO

import pytest

from neurokiln.report import (
    _tex,
    build_docx,
    build_html_report,
    build_latex,
    build_workbook,
)


@pytest.fixture()
def detail():
    return {
        "id": 42,
        "name": "CSP vs Riemann 对照",
        "status": "completed",
        "protocol": "within_subject",
        "protocol_params": {"folds": 2, "repeats": 1, "seeds": [0]},
        "notes": "回归样例",
        "engine_version": "0.4.0",
        "duration_ms": 1234.0,
        "seeds": [0],
        "n_runs": 2,
        "dataset": {
            "id": "synthetic",
            "name": "合成数据",
            "version": "1",
            "signature": "sig_abc-123",
            "manifest": {
                "subjects": ["S01", "S02"],
                "loaded_n_trials": 40,
                "loaded_channels": 8,
                "loaded_sfreq": 128,
                "label_names": ["左手", "右手"],
                "citation": "NeuroKiln synthetic generator",
            },
        },
        "rows": [
            {
                "pipeline": "csp_lda",
                "pipeline_name": "CSP + LDA",
                "balanced_accuracy_mean": 0.81,
                "balanced_accuracy_std": 0.02,
                "balanced_accuracy_ci95": [0.78, 0.84],
                "cohen_kappa_mean": 0.62,
                "calibration": {"ece": 0.05},
                "splits": [
                    {
                        "duration_ms": 12.5,
                        "y_true": [0, 1, 0, 1],
                        "y_pred": [0, 1, 0, 0],
                    }
                ],
            },
            {
                "pipeline": "riemann_mdm",
                "pipeline_name": "Riemann MDM",
                "balanced_accuracy_mean": 0.74,
                "balanced_accuracy_std": 0.03,
                "cohen_kappa_mean": 0.48,
                "calibration": {"ece": 0.08},
                "splits": [{"duration_ms": 9.0, "y_true": [0, 1], "y_pred": [0, 1]}],
            },
        ],
        "statistics": {
            "unit": "subject",
            "n_units": 2,
            "n_splits": 2,
            "alpha": 0.05,
            "friedman": {"available": True, "statistic": 3.0, "p_value": 0.083,
                         "significant": False},
            "nemenyi": {"available": True, "critical_difference": 1.96},
            "pairwise": [
                {"pipeline_a": "csp_lda", "pipeline_b": "riemann_mdm", "statistic": 1.0,
                 "p_value": 0.5, "p_value_adjusted": 0.5, "significant": False}
            ],
        },
    }


def test_latex_report_escapes_and_lists_gbt_references(detail):
    text = build_latex(detail)
    assert text.startswith("% 使用 XeLaTeX 编译")
    assert r"\begin{document}" in text and r"\end{document}" in text
    assert r"\usepackage{booktabs}" in text
    assert r"95\% CI" in text
    assert r"sig\_abc-123" in text
    assert "CSP + LDA" in text
    assert r"JAYARAM V, BARACHANT A" in text
    assert r"\bibitem" in text
    assert "NeuroKiln synthetic generator" in text
    assert "平衡准确率" in text


def test_docx_report_contains_title_tables_and_references(detail):
    payload = build_docx(detail)
    assert payload[:2] == b"PK"
    from docx import Document

    document = Document(BytesIO(payload))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "NeuroKiln 实验报告" in text
    assert "参考文献" in text
    assert "JAYARAM V, BARACHANT A" in text
    assert len(document.tables) >= 3
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "CSP + LDA" in table_text
    assert "0.8100" in table_text


def test_tex_escaping_handles_special_characters():
    assert _tex("a_b & 50% #1 {x} ~ ^ \\") == (
        r"a\_b \& 50\% \#1 \{x\} \textasciitilde{} \textasciicircum{} \textbackslash{}"
    )


def _explanation_item():
    return {
        "id": 3,
        "pipeline": "csp_lda",
        "pipeline_name": "CSP + LDA",
        "fold": 0,
        "repeat": 0,
        "seed": 0,
        "status": "completed",
        "summary": {
            "performance": {"balanced_accuracy": 0.81},
            "warnings": ["测试警告"],
        },
        "document": {
            "dataset": {"channel_names": ["C3", "C4"]},
            "faithfulness": {
                "deletion_auc": 0.62,
                "insertion_auc": 1.18,
                "random_deletion_auc": 1.01,
                "aopc": 0.12,
                "normalization": {"usable": True},
                "stability": {"half_split_spearman": 0.4},
            },
            "audit": {
                "closure_ratio": 0.95,
                "channel_importance": [0.3, 0.1],
                "topomap": "audit.png",
            },
            "occlusion": {
                "channel_importance": [0.4, -0.2],
                "topomap": "occlusion.png",
            },
            "bands": {
                "topomap": "band.png",
                "topomap_band": [8, 13],
                "channel_importance": {"8-13": [0.5, -0.1]},
            },
            "gradient": {"available": False},
            "intrinsic": [
                {
                    "id": "csp_patterns",
                    "label": "CSP 空间模式 × LDA 权重",
                    "values": [1.0, -1.0],
                    "topomap": "csp.png",
                }
            ],
        },
        "images": {
            "occlusion.png": "data:image/png;base64,AAAA",
            "csp.png": "data:image/png;base64,BBBB",
        },
    }


def test_html_report_renders_explanation_section(detail):
    detail = {**detail, "explanations": [_explanation_item()]}
    html = build_html_report(detail)
    assert "模型在看哪里（可解释性）" in html
    assert "CSP + LDA · fold 0" in html
    assert "删除 AUC" in html and "0.6200" in html
    assert "测试警告" in html
    assert "data:image/png;base64,AAAA" in html
    assert "C3" in html and "0.4000" in html

    empty = build_html_report({**detail, "explanations": []})
    assert "本次实验尚未生成可解释性分析" in empty


def test_workbook_contains_explanation_sheets(detail):
    from openpyxl import load_workbook

    detail = {**detail, "explanations": [_explanation_item()]}
    book = load_workbook(BytesIO(build_workbook(detail)))
    assert "模型在看哪里" in book.sheetnames
    assert "解释通道重要性" in book.sheetnames
    rows = list(book["模型在看哪里"].iter_rows(min_row=2, values_only=True))
    assert rows and rows[0][1] == "CSP + LDA"
    assert rows[0][6] == 0.62
    channel_rows = list(book["解释通道重要性"].iter_rows(min_row=2, values_only=True))
    assert channel_rows and channel_rows[0][2] == "C3"
