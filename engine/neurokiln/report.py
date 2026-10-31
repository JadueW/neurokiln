from __future__ import annotations

import base64
import html
from datetime import datetime
from io import BytesIO
from typing import Any, Sequence

import numpy as np

from .evaluate import confusion_matrix

_REPORT_CSS = """
:root { color-scheme: light; }
* { box-sizing: border-box; }
body { margin: 0; padding: 32px; background: #eef2f6; color: #1f2933;
  font-family: "Segoe UI", "PingFang SC", "Microsoft YaHei", sans-serif; line-height: 1.6; }
main { max-width: 1180px; margin: 0 auto; background: #fff; border-radius: 14px;
  box-shadow: 0 18px 50px rgba(31, 45, 61, .12); overflow: hidden; }
header { padding: 34px 42px 28px; color: #fff; background:
  linear-gradient(135deg, #1d2a44 0%, #223a63 52%, #14538a 100%); }
header .eyebrow { letter-spacing: .18em; font-size: 12px; opacity: .68; text-transform: uppercase; }
header h1 { margin: 8px 0 10px; font-size: 27px; font-weight: 650; }
header .meta { display: flex; flex-wrap: wrap; gap: 8px 26px; font-size: 13px; opacity: .86; }
section { padding: 26px 42px; border-bottom: 1px solid #e8edf3; }
section:last-child { border-bottom: 0; }
h2 { margin: 0 0 16px; font-size: 18px; color: #16233a; display: flex; align-items: center; gap: 9px; }
h2::before { content: ""; width: 4px; height: 18px; border-radius: 2px; background: #2f6fb1; }
h3 { margin: 20px 0 10px; font-size: 14px; color: #31425c; }
.cards { display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr)); gap: 12px; }
.card { padding: 15px 18px; border: 1px solid #e3eaf3; border-radius: 10px; background: #f8fafc; }
.card .label { font-size: 12px; color: #6b7c93; }
.card .value { margin-top: 4px; font-size: 21px; font-weight: 650; color: #1c3d68; }
.card .hint { margin-top: 2px; font-size: 11px; color: #8a97a8; word-break: break-all; }
table { width: 100%; border-collapse: collapse; font-size: 12.5px; }
th, td { padding: 9px 11px; text-align: left; border-bottom: 1px solid #edf1f6; vertical-align: top; }
th { background: #f4f7fb; color: #50617a; font-weight: 600; white-space: nowrap; }
tbody tr:hover { background: #fafcff; }
td.num, th.num { text-align: right; font-variant-numeric: tabular-nums; }
.mono { font-family: "Cascadia Mono", Consolas, monospace; font-size: 11.5px; }
.tag { display: inline-block; padding: 1px 8px; border-radius: 10px; font-size: 11px; background: #e9f0fa; color: #2b5f9e; }
.tag.bad { background: #fdecec; color: #b23c3c; }
.tag.good { background: #e8f7ef; color: #217a4b; }
.note { margin: 10px 0 0; padding: 10px 14px; border-radius: 8px; background: #f4f8fd;
  border-left: 3px solid #4c86c8; color: #40536c; font-size: 12.5px; }
.bar { display: flex; align-items: center; gap: 10px; margin: 7px 0; }
.bar .name { width: 150px; font-size: 12.5px; color: #35475f; overflow: hidden; text-overflow: ellipsis; white-space: nowrap; }
.bar .track { flex: 1; height: 13px; border-radius: 7px; background: #eef2f7; overflow: hidden; }
.bar .fill { height: 100%; border-radius: 7px; background: linear-gradient(90deg, #4d8bcb, #2f6fb1); }
.bar .value { width: 56px; text-align: right; font-size: 12px; font-variant-numeric: tabular-nums; }
.cm { display: inline-block; margin: 0 18px 18px 0; vertical-align: top; }
.cm table { width: auto; }
.cm td, .cm th { text-align: center; min-width: 44px; padding: 6px 8px; }
.cm td.hot { color: #fff; }
.xai-grid { display: grid; grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
  gap: 14px; margin-top: 16px; }
.xai-fig { margin: 0; padding: 10px; border: 1px solid #e3eaf3; border-radius: 10px;
  background: #fbfdff; text-align: center; }
.xai-fig img { width: 100%; max-width: 240px; image-rendering: auto; }
.xai-fig figcaption { margin-top: 6px; font-size: 11.5px; color: #5b6b80; }
.xai-warn { margin: 10px 0 0; padding-left: 20px; color: #8a5a12; font-size: 12.5px; }
details { margin-top: 10px; }
summary { cursor: pointer; color: #2f6fb1; font-size: 12.5px; }
footer { padding: 20px 42px 28px; color: #8a97a8; font-size: 11.5px; background: #fafcfe; }
@media print { body { padding: 0; background: #fff; } main { box-shadow: none; border-radius: 0; } section { break-inside: avoid; } }
"""


def _e(value: Any) -> str:
    return html.escape("" if value is None else str(value))


def _num(value: Any, digits: int = 4) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "—"
    if not np.isfinite(number):
        return "—"
    return f"{number:.{digits}f}"


def _p_value(value: Any) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "—"
    if not np.isfinite(number):
        return "—"
    return f"{number:.2e}" if number < 0.001 else f"{number:.4f}"


def _ci(value: Any) -> str:
    if not isinstance(value, (list, tuple)) or len(value) != 2:
        return "—"
    if value[0] is None or value[1] is None:
        return "—"
    return f"[{float(value[0]):.3f}, {float(value[1]):.3f}]"


def _metric_cards(detail: dict[str, Any]) -> str:
    summary = detail.get("summary") or {}
    rows = detail.get("rows") or []
    best = max(rows, key=lambda row: row.get("balanced_accuracy_mean", 0.0), default=None)
    stats = detail.get("statistics") or summary.get("statistics") or {}
    pairwise = stats.get("pairwise") or []
    significant = sum(1 for item in pairwise if item.get("significant"))
    cards = [
        ("最佳算法", (best or {}).get("pipeline_name") or (best or {}).get("pipeline") or "—",
         f"BA {_num((best or {}).get('balanced_accuracy_mean'))}"),
        ("检验单位 / 样本量",
         f"{'被试' if stats.get('unit') == 'subject' else '折'} {stats.get('n_units', 0)}",
         f"{stats.get('n_splits', 0)} 折 · {detail.get('n_runs', 0)} 次运行 · 种子 {detail.get('seeds') or []}"),
        ("Friedman p", _p_value((stats.get("friedman") or {}).get("p_value")),
         "显著" if (stats.get("friedman") or {}).get("significant") else "未达显著"),
        ("显著算法对", str(significant), f"Holm 校正，α={stats.get('alpha', 0.05)}"),
        ("总耗时", f"{_num(float(detail.get('duration_ms') or 0) / 1000.0, 1)} s",
         f"引擎 {detail.get('engine_version', '—')}"),
    ]
    cells = "".join(
        f'<div class="card"><div class="label">{_e(label)}</div>'
        f'<div class="value">{_e(value)}</div><div class="hint">{_e(hint)}</div></div>'
        for label, value, hint in cards
    )
    return f'<div class="cards">{cells}</div>'


def _summary_table(rows: Sequence[dict[str, Any]]) -> str:
    ordered = sorted(rows, key=lambda row: row.get("balanced_accuracy_mean", 0.0), reverse=True)
    body = []
    for rank, row in enumerate(ordered, start=1):
        splits = row.get("splits") or []
        durations = [float(item.get("duration_ms", 0.0)) for item in splits]
        body.append(
            "<tr>"
            f'<td class="num">{rank}</td>'
            f"<td><strong>{_e(row.get('pipeline_name') or row.get('pipeline'))}</strong>"
            f'<div class="mono">{_e(row.get("pipeline"))}</div></td>'
            f'<td class="num">{_num(row.get("balanced_accuracy_mean"))}</td>'
            f'<td class="num">± {_num(row.get("balanced_accuracy_std"))}</td>'
            f'<td class="num">{_ci(row.get("balanced_accuracy_ci95"))}</td>'
            f'<td class="num">{_num(row.get("cohen_kappa_mean"))}</td>'
            f'<td class="num">{_num((row.get("calibration") or {}).get("ece"))}</td>'
            f'<td class="num">{_num(np.mean(durations) if durations else 0.0, 1)} ms</td>'
            f'<td class="num">{len(splits)}</td>'
            "</tr>"
        )
    return (
        '<table><thead><tr><th class="num">#</th><th>算法</th>'
        '<th class="num">平衡准确率</th><th class="num">标准差</th>'
        '<th class="num">95% CI</th><th class="num">Cohen κ</th>'
        '<th class="num">ECE</th>'
        '<th class="num">平均单折耗时</th><th class="num">折数</th>'
        '</tr></thead><tbody>'
        + "".join(body)
        + "</tbody></table>"
    )


def _chart(rows: Sequence[dict[str, Any]]) -> str:
    ordered = sorted(rows, key=lambda row: row.get("balanced_accuracy_mean", 0.0), reverse=True)
    bars = []
    for row in ordered:
        value = float(row.get("balanced_accuracy_mean", 0.0) or 0.0)
        width = max(0.0, min(100.0, value * 100.0))
        bars.append(
            f'<div class="bar"><div class="name">{_e(row.get("pipeline_name") or row.get("pipeline"))}</div>'
            f'<div class="track"><div class="fill" style="width:{width:.1f}%"></div></div>'
            f'<div class="value">{value:.3f}</div></div>'
        )
    return "".join(bars)


def _statistics_section(detail: dict[str, Any]) -> str:
    stats = detail.get("statistics") or (detail.get("summary") or {}).get("statistics") or {}
    if not stats or not stats.get("n_splits"):
        return '<p class="note">该实验折数不足，未生成显著性检验。</p>'
    friedman = stats.get("friedman") or {}
    nemenyi = stats.get("nemenyi") or {}
    unit_label = "被试" if stats.get("unit") == "subject" else "折"
    parts = [
        "<h2>统计检验</h2>",
        f'<p class="note">{_e(stats.get("note", ""))} Friedman 检验跨算法比较同一{unit_label}上的平衡准确率；'
        "Nemenyi 临界距离基于 Demsar (2006) 的平均秩方法。"
        + (
            " 被试级检验用于控制折叠间相关性；单被试实验仍会回退到折级并明确标注。"
            if stats.get("unit") == "subject" else ""
        )
        + "</p>",
    ]
    if friedman.get("available"):
        verdict = "差异显著" if friedman.get("significant") else "差异未达显著"
        parts.append(
            f'<p>Friedman χ² = <strong>{_num(friedman.get("statistic"), 3)}</strong>，'
            f'p = <strong>{_p_value(friedman.get("p_value"))}</strong>'
            f'<span class="tag {"good" if friedman.get("significant") else ""}">{verdict}</span></p>'
        )
    else:
        parts.append(f'<p class="note">{_e(friedman.get("reason") or "Friedman 不可用")}</p>')

    ranks = (stats.get("average_ranks") or {}).get("ranks") or []
    if ranks:
        cells = "".join(
            f"<tr><td>{_e(item['pipeline'])}</td><td class=\"num\">{_num(item['mean_rank'], 2)}</td>"
            f"<td class=\"num\">{_num(item['mean_balanced_accuracy'])}</td></tr>"
            for item in ranks
        )
        parts.append(
            "<h3>平均秩（秩越小越好）</h3>"
            '<table><thead><tr><th>算法</th><th class="num">平均秩</th>'
            f'<th class="num">平均 BA</th></tr></thead><tbody>{cells}</tbody></table>'
        )
        if nemenyi.get("available"):
            parts.append(
                f'<p class="note">Nemenyi 临界距离 CD = <strong>{_num(nemenyi.get("critical_difference"), 3)}</strong>'
                f'（α={stats.get("alpha", 0.05)}，{stats.get("n_splits")} 折）。</p>'
            )

    pairwise = stats.get("pairwise") or []
    if pairwise:
        cells = []
        for item in pairwise:
            verdict = (
                '<span class="tag good">显著</span>'
                if item.get("significant")
                else '<span class="tag">不显著</span>'
            )
            cells.append(
                "<tr>"
                f"<td>{_e(item['pipeline_a'])} vs {_e(item['pipeline_b'])}</td>"
                f"<td class=\"num\">{_num(item.get('statistic'), 3)}</td>"
                f"<td class=\"num\">{_p_value(item.get('p_value'))}</td>"
                f"<td class=\"num\">{_p_value(item.get('p_value_adjusted'))}</td>"
                f"<td>{verdict}</td>"
                "</tr>"
            )
        parts.append(
            "<h3>Wilcoxon 符号秩检验（Holm 校正）</h3>"
            '<table><thead><tr><th>算法对</th><th class="num">统计量</th>'
            '<th class="num">原始 p</th><th class="num">校正 p</th><th>结论</th>'
            f"</tr></thead><tbody>{''.join(cells)}</tbody></table>"
        )
    return "".join(parts)


def _calibration_section(detail: dict[str, Any]) -> str:
    rows = detail.get("rows") or []
    available = [row for row in rows if (row.get("calibration") or {}).get("available")]
    if not available:
        return (
            '<p class="note">本次实验没有可用的 predict_proba 概率输出，'
            "无法计算校准指标；请检查插件或自定义解码器是否实现了 predict_proba。</p>"
        )
    cards = "".join(
        f'<div class="card"><div class="label">{_e(row.get("pipeline_name") or row.get("pipeline"))}</div>'
        f'<div class="value">ECE {_num((row.get("calibration") or {}).get("ece"))}</div>'
        f'<div class="hint">Brier {_num((row.get("calibration") or {}).get("brier"))} · '
        f'NLL {_num((row.get("calibration") or {}).get("nll"))} · '
        f'{_e((row.get("calibration") or {}).get("n_evaluated"))} 试次</div></div>'
        for row in available
    )
    tables = []
    for row in available:
        calibration = row.get("calibration") or {}
        curve = calibration.get("reliability") or []
        cells = "".join(
            "<tr>"
            f'<td class="num">{_e(item.get("bin_lower"))}–{_e(item.get("bin_upper"))}</td>'
            f'<td class="num">{_num(item.get("confidence"))}</td>'
            f'<td class="num">{_num(item.get("accuracy"))}</td>'
            f'<td class="num">{_e(item.get("count"))}</td>'
            "</tr>"
            for item in curve
        )
        tables.append(
            f'<h3>{_e(row.get("pipeline_name") or row.get("pipeline"))}</h3>'
            '<table><thead><tr><th class="num">置信度区间</th>'
            '<th class="num">平均置信度</th><th class="num">实际正确率</th>'
            f'<th class="num">试次数</th></tr></thead><tbody>{cells}</tbody></table>'
        )
    return (
        '<p class="note">校准指标基于各折留出预测的概率拼接计算：ECE 为按样本加权的 '
        "置信度–正确率偏差，Brier 为多分类平方误差，NLL 为真实类别负对数似然。"
        "概率来自训练折拟合的模型，测试折只参与评估。</p>"
        f'<div class="cards">{cards}</div>'
        + "".join(tables)
    )


def _confusion_matrices(detail: dict[str, Any]) -> str:
    label_names = (
        (detail.get("dataset") or {}).get("manifest", {}).get("label_names") or []
    )
    blocks = []
    for row in detail.get("rows") or []:
        splits = row.get("splits") or []
        y_true = [label for split in splits for label in split.get("y_true", [])]
        y_pred = [label for split in splits for label in split.get("y_pred", [])]
        if not y_true or not y_pred or len(y_true) != len(y_pred):
            continue
        classes = sorted(set(y_true) | set(y_pred))
        matrix = confusion_matrix(np.asarray(y_true), np.asarray(y_pred), len(classes))
        labels = [
            label_names[index] if index < len(label_names) else f"类别 {index}"
            for index in classes
        ]
        head = "".join(f"<th>{_e(label)}</th>" for label in labels)
        body = []
        for row_index, label in enumerate(labels):
            row_total = max(1, int(matrix[row_index].sum()))
            cells = []
            for column_index in range(len(classes)):
                value = int(matrix[row_index, column_index])
                ratio = value / row_total
                if ratio > 0.65:
                    style, klass = "background:#2f6fb1;color:#fff", "hot"
                elif ratio > 0.35:
                    style, klass = f"background:rgba(47,111,177,{0.12 + ratio * 0.55:.2f})", ""
                else:
                    style, klass = f"background:rgba(47,111,177,{ratio * 0.22:.2f})", ""
                cells.append(f'<td class="{klass}" style="{style}">{value}</td>')
            body.append(f"<tr><th>{_e(label)}</th>{''.join(cells)}</tr>")
        blocks.append(
            f'<div class="cm"><h3>{_e(row.get("pipeline_name") or row.get("pipeline"))}</h3>'
            f'<table><thead><tr><th>真实 \\ 预测</th>{head}</tr></thead><tbody>{"".join(body)}</tbody></table></div>'
        )
    return "".join(blocks) or '<p class="note">没有保存逐试次预测。</p>'


def _fold_table(row: dict[str, Any]) -> str:
    cells = "".join(
        "<tr>"
        f"<td class=\"num\">{_e((split.get('split') or {}).get('fold', '—'))}</td>"
        f"<td class=\"num\">{_e((split.get('split') or {}).get('repeat', '—'))}</td>"
        f"<td class=\"num\">{_e((split.get('split') or {}).get('seed', '—'))}</td>"
        f"<td class=\"num\">{_e((split.get('split') or {}).get('n_train', '—'))}</td>"
        f"<td class=\"num\">{_e((split.get('split') or {}).get('n_test', '—'))}</td>"
        f"<td class=\"num\">{_num(split.get('balanced_accuracy'))}</td>"
        f"<td class=\"num\">{_num(split.get('cohen_kappa'))}</td>"
        f"<td class=\"num\">{_num(split.get('duration_ms'), 1)} ms</td>"
        "</tr>"
        for split in row.get("splits") or []
    )
    return (
        '<table><thead><tr><th class="num">Fold</th><th class="num">Repeat</th>'
        '<th class="num">Seed</th><th class="num">训练试次</th><th class="num">测试试次</th>'
        '<th class="num">平衡准确率</th><th class="num">Cohen κ</th><th class="num">耗时</th>'
        f"</tr></thead><tbody>{cells}</tbody></table>"
    )


def _explanation_items(detail: dict[str, Any]) -> list[dict[str, Any]]:
    items = []
    for item in detail.get("explanations") or []:
        if isinstance(item, dict) and (item.get("document") or item.get("summary")):
            items.append(item)
    return items


def _xai_label(item: dict[str, Any]) -> str:
    name = item.get("pipeline_name") or item.get("pipeline") or "算法"
    return f"{name} · fold {item.get('fold', 0)}"


def _xai_metrics(item: dict[str, Any]) -> list[tuple[str, str, str]]:
    document = item.get("document") or {}
    summary = item.get("summary") or {}
    performance = summary.get("performance") or {}
    faith = document.get("faithfulness") or {}
    norm = faith.get("normalization") or {}
    audit = document.get("audit") or {}
    stability = (faith.get("stability") or {}).get("half_split_spearman")
    usable = norm.get("usable", True)
    return [
        ("平衡准确率", _num(performance.get("balanced_accuracy")), "该折留出预测"),
        (
            "删除 AUC",
            _num(faith.get("deletion_auc")),
            "越低越忠实" if usable else "归一化不可用",
        ),
        (
            "插入 AUC",
            _num(faith.get("insertion_auc")),
            "越高越忠实" if usable else "归一化不可用",
        ),
        ("随机删除 AUC", _num(faith.get("random_deletion_auc")), "对照基线"),
        ("AOPC", _num(faith.get("aopc")), "前 8 步平均下降"),
        (
            "closure",
            _num(audit.get("closure_ratio")),
            "已知频带/Hjorth 可解释比例",
        ),
        (
            "半折稳定性 ρ",
            _num(stability),
            "两半测试集通道排序一致性",
        ),
    ]


def _xai_channel_rows(
    item: dict[str, Any], limit: int = 12, formatted: bool = True
) -> list[list[Any]]:
    document = item.get("document") or {}
    channels = (document.get("dataset") or {}).get("channel_names") or []
    occlusion = document.get("occlusion") or {}
    values = occlusion.get("channel_importance") or []
    if not channels or len(values) != len(channels):
        return []
    intrinsic = (document.get("intrinsic") or [{}])[0] if document.get("intrinsic") else {}
    intrinsic_values = intrinsic.get("values") or []
    bands = document.get("bands") or {}
    band_values = (bands.get("channel_importance") or {}).get(
        "-".join(f"{value:g}" for value in (bands.get("topomap_band") or []))
    ) or []
    gradient = document.get("gradient") or {}
    gradient_values = gradient.get("channel_magnitude") or []
    audit = document.get("audit") or {}
    audit_values = audit.get("channel_importance") or []

    def value_at(collection, index):
        if index >= len(collection) or collection[index] is None:
            return "—" if formatted else None
        return _num(collection[index], 4) if formatted else float(collection[index])

    order = sorted(range(len(channels)), key=lambda index: -abs(values[index]))[:limit]
    return [
        [
            channels[index],
            _num(values[index], 4) if formatted else float(values[index]),
            value_at(intrinsic_values, index),
            value_at(band_values, index),
            value_at(gradient_values, index),
            value_at(audit_values, index),
        ]
        for index in order
    ]


def _xai_figures(item: dict[str, Any]) -> str:
    document = item.get("document") or {}
    images = item.get("images") or {}
    figures = []

    def figure(asset, caption):
        if not asset or asset not in images:
            return
        figures.append(
            f'<figure class="xai-fig"><img src="{images[asset]}" alt="{_e(caption)}">'
            f"<figcaption>{_e(caption)}</figcaption></figure>"
        )

    figure((document.get("occlusion") or {}).get("topomap"), "遮挡归因（Δlog-odds）")
    bands = document.get("bands") or {}
    if bands.get("topomap"):
        band_label = " / ".join(f"{value:g} Hz" for value in bands.get("topomap_band") or [])
        figure(bands["topomap"], f"频带消融（{band_label}）")
    figure((document.get("gradient") or {}).get("topomap"), "梯度归因")
    for entry in (document.get("intrinsic") or [])[:2]:
        figure(entry.get("topomap"), f"内在参数：{entry.get('label') or entry.get('id')}")
    figure((document.get("audit") or {}).get("topomap"), "已知特征（频带能量）重要性")
    if not figures:
        return ""
    return f'<div class="xai-grid">{"".join(figures)}</div>'


def _explanations_section(detail: dict[str, Any]) -> str:
    items = _explanation_items(detail)
    if not items:
        return (
            '<p class="note">本次实验尚未生成可解释性分析。在「可解释性」页选择实验、'
            "算法与折后运行，结果会自动出现在本报告中。</p>"
        )
    blocks = []
    for item in items:
        document = item.get("document") or {}
        summary = item.get("summary") or {}
        cards = "".join(
            f'<div class="card"><div class="label">{_e(label)}</div>'
            f'<div class="value">{_e(value)}</div><div class="hint">{_e(hint)}</div></div>'
            for label, value, hint in _xai_metrics(item)
        )
        channel_rows = "".join(
            "<tr>"
            + "".join(f'<td class="num">{_e(cell)}</td>' for cell in row)
            + "</tr>"
            for row in _xai_channel_rows(item)
        )
        channel_table = (
            '<table><thead><tr><th>通道</th><th class="num">遮挡 Δ</th>'
            '<th class="num">内在参数</th><th class="num">频带消融</th>'
            '<th class="num">梯度</th><th class="num">频带能量</th></tr></thead>'
            f"<tbody>{channel_rows}</tbody></table>"
            if channel_rows
            else '<p class="note">该算法没有可映射到电极的通道重要性。</p>'
        )
        warnings = summary.get("warnings") or []
        warning_html = (
            '<ul class="xai-warn">'
            + "".join(f"<li>{_e(message)}</li>" for message in warnings)
            + "</ul>"
            if warnings
            else ""
        )
        faith = document.get("faithfulness") or {}
        agreement = faith.get("agreement") or {}
        agreement_bits = []
        intrinsic_agreement = (agreement.get("intrinsic") or {}).get("spearman")
        if intrinsic_agreement is not None:
            agreement_bits.append(f"遮挡↔内在参数 ρ={float(intrinsic_agreement):.3f}")
        gradient_agreement = (agreement.get("gradient") or {}).get("spearman")
        if gradient_agreement is not None:
            agreement_bits.append(f"遮挡↔梯度 ρ={float(gradient_agreement):.3f}")
        sanity = document.get("sanity") or {}
        if sanity.get("available") and sanity.get("spearman") is not None:
            agreement_bits.append(f"权重随机化 ρ={float(sanity['spearman']):.3f}")
        agreement_note = (
            '<p class="note">' + _e("；".join(agreement_bits)) + "。</p>"
            if agreement_bits
            else ""
        )
        blocks.append(
            f"<h3>{_e(_xai_label(item))}</h3>"
            f'<div class="cards">{cards}</div>'
            f"{_xai_figures(item)}"
            f'<div style="margin-top:12px">{channel_table}</div>'
            f"{agreement_note}{warning_html}"
        )
    prov = (items[0].get("document") or {}).get("provenance") or {}
    note = (
        "遮挡基线为训练折逐通道均值，测试折不参与任何拟合；"
        "归因目标是模型自身预测类别的 log-odds。"
        f"背景来源 {_e(prov.get('background', 'train_fold_only'))}。"
    )
    return f'<p class="note">{note}</p>' + "".join(blocks)


def build_html_report(detail: dict[str, Any]) -> str:
    dataset = detail.get("dataset") or {}
    manifest = dataset.get("manifest") or {}
    protocol_params = detail.get("protocol_params") or {}
    rows = detail.get("rows") or []
    pipelines = detail.get("pipelines") or []
    pipeline_names = {item.get("id"): item.get("name") for item in pipelines}
    for row in rows:
        if not row.get("pipeline_name"):
            row["pipeline_name"] = pipeline_names.get(row.get("pipeline"), row.get("pipeline"))

    details_blocks = []
    for row in sorted(rows, key=lambda item: item.get("balanced_accuracy_mean", 0.0), reverse=True):
        details_blocks.append(
            f'<details><summary>{_e(row.get("pipeline_name") or row.get("pipeline"))} · 逐折明细</summary>'
            f"{_fold_table(row)}</details>"
        )

    generated = datetime.now().strftime("%Y-%m-%d %H:%M")
    title = f"NeuroKiln · 实验 #{detail.get('id')} 报告"
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{_e(title)}</title>
<style>{_REPORT_CSS}</style>
</head>
<body>
<main>
<header>
  <div class="eyebrow">NeuroKiln · Brain–Computer Interface Algorithm Evaluation</div>
  <h1>{_e(detail.get("name") or title)}</h1>
  <div class="meta">
    <span>实验 #{_e(detail.get("id"))}</span>
    <span>数据集 {_e(dataset.get("name") or dataset.get("id"))} · v{_e(dataset.get("version"))}</span>
    <span>协议 {_e(detail.get("protocol"))}</span>
    <span>状态 {_e(detail.get("status"))}</span>
    <span>生成于 {_e(generated)}</span>
  </div>
</header>
<section>
  {_metric_cards(detail)}
  <p class="note">所有预处理统计量、特征提取与分类器参数仅在训练折拟合；预测来自留出折。数据签名
  <span class="mono">{_e(dataset.get("signature"))}</span> 用于识别数据集版本。</p>
</section>
<section>
  <h2>算法排名</h2>
  {_chart(rows)}
  <div style="margin-top:14px">{_summary_table(rows)}</div>
</section>
<section>
  {_statistics_section(detail)}
</section>
<section>
  <h2>概率校准</h2>
  {_calibration_section(detail)}
</section>
<section>
  <h2>混淆矩阵（全折汇总）</h2>
  {_confusion_matrices(detail)}
</section>
<section>
  <h2>模型在看哪里（可解释性）</h2>
  {_explanations_section(detail)}
</section>
<section>
  <h2>实验配置</h2>
  <table>
    <tbody>
      <tr><th>数据集</th><td>{_e(dataset.get("name"))} <span class="mono">{_e(dataset.get("signature"))}</span></td></tr>
      <tr><th>被试</th><td>{_e(", ".join(str(item) for item in (manifest.get("subjects") or [])) or "—")}</td></tr>
      <tr><th>试次 / 通道 / 采样率</th><td>{_e(manifest.get("loaded_n_trials"))} / {_e(manifest.get("loaded_channels"))} / {_e(manifest.get("loaded_sfreq"))} Hz</td></tr>
      <tr><th>协议参数</th><td class="mono">{_e(protocol_params)}</td></tr>
      <tr><th>备注</th><td>{_e(detail.get("notes") or "—")}</td></tr>
    </tbody>
  </table>
</section>
<section>
  <h2>逐折明细</h2>
  {"".join(details_blocks) or '<p class="note">没有逐折记录。</p>'}
</section>
<footer>
  由 NeuroKiln 自动生成 · 引擎 {_e(detail.get("engine_version"))} · 本报告不包含可识别被试身份的信息。
</footer>
</main>
</body>
</html>"""


def _safe(value: Any) -> Any:
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def build_workbook(detail: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1F3A5F")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    title_font = Font(color="1F3A5F", bold=True, size=14)
    fine = Side(style="thin", color="DCE4EF")
    border = Border(left=fine, right=fine, top=fine, bottom=fine)

    def style_header(sheet, row: int, columns: int) -> None:
        for column in range(1, columns + 1):
            cell = sheet.cell(row=row, column=column)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    def autosize(sheet, widths: dict[int, float]) -> None:
        for index, width in widths.items():
            sheet.column_dimensions[get_column_letter(index)].width = width

    dataset = detail.get("dataset") or {}
    manifest = dataset.get("manifest") or {}
    stats = detail.get("statistics") or (detail.get("summary") or {}).get("statistics") or {}
    rows = sorted(
        detail.get("rows") or [],
        key=lambda item: item.get("balanced_accuracy_mean", 0.0),
        reverse=True,
    )

    book = Workbook()
    overview = book.active
    overview.title = "概览"
    overview["A1"] = "NeuroKiln 实验报告"
    overview["A1"].font = title_font
    info = [
        ("实验 ID", detail.get("id")),
        ("实验名称", detail.get("name")),
        ("数据集", dataset.get("name") or dataset.get("id")),
        ("数据集版本", dataset.get("version")),
        ("数据签名", dataset.get("signature")),
        ("协议", detail.get("protocol")),
        ("协议参数", str(detail.get("protocol_params") or {})),
        ("状态", detail.get("status")),
        ("引擎版本", detail.get("engine_version")),
        ("折数", stats.get("n_splits")),
        ("检验单位", stats.get("unit")),
        ("检验样本量", stats.get("n_units")),
        ("被试数", stats.get("n_subjects")),
        ("运行数", detail.get("n_runs")),
        ("总耗时 (s)", round(float(detail.get("duration_ms") or 0.0) / 1000.0, 2)),
        ("被试", ", ".join(str(item) for item in (manifest.get("subjects") or []))),
        ("试次 / 通道 / 采样率", f"{manifest.get('loaded_n_trials')} / {manifest.get('loaded_channels')} / {manifest.get('loaded_sfreq')} Hz"),
        ("备注", detail.get("notes") or ""),
    ]
    for offset, (label, value) in enumerate(info, start=3):
        overview.cell(row=offset, column=1, value=label).font = Font(bold=True, color="40536C")
        overview.cell(row=offset, column=2, value=_safe(str(value) if value is not None else "—"))
        overview.cell(row=offset, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    overview.column_dimensions["A"].width = 24
    overview.column_dimensions["B"].width = 88

    summary = book.create_sheet("算法汇总")
    headers = ["排名", "算法 ID", "算法名称", "平衡准确率均值", "标准差", "95% CI",
               "Cohen κ 均值", "ECE", "Brier", "NLL", "平均单折耗时 (ms)", "折数"]
    summary.append(headers)
    style_header(summary, 1, len(headers))
    for rank, row in enumerate(rows, start=1):
        splits = row.get("splits") or []
        calibration = row.get("calibration") or {}
        summary.append([
            rank,
            row.get("pipeline"),
            row.get("pipeline_name") or row.get("pipeline"),
            round(float(row.get("balanced_accuracy_mean") or 0.0), 6),
            round(float(row.get("balanced_accuracy_std") or 0.0), 6),
            _ci(row.get("balanced_accuracy_ci95")),
            round(float(row.get("cohen_kappa_mean") or 0.0), 6),
            calibration.get("ece"),
            calibration.get("brier"),
            calibration.get("nll"),
            round(float(np.mean([item.get("duration_ms", 0.0) for item in splits])) if splits else 0.0, 2),
            len(splits),
        ])
    for row in summary.iter_rows(min_row=2, max_row=summary.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
        row[3].number_format = "0.0000"
        row[4].number_format = "0.0000"
        row[6].number_format = "0.0000"
        for index in (7, 8, 9):
            row[index].number_format = "0.0000"
    if summary.max_row > 1:
        summary.conditional_formatting.add(
            f"D2:D{summary.max_row}",
            ColorScaleRule(start_type="min", start_color="F8D7DA",
                           mid_type="percentile", mid_value=50, mid_color="FFF2CC",
                           end_type="max", end_color="C6E7D4"),
        )
    autosize(summary, {1: 8, 2: 20, 3: 30, 4: 18, 5: 12, 6: 16, 7: 14, 8: 10, 9: 10, 10: 10, 11: 20, 12: 8})
    summary.freeze_panes = "A2"

    folds = book.create_sheet("逐折结果")
    fold_headers = ["算法 ID", "算法名称", "Fold", "Repeat", "Seed", "训练试次", "测试试次",
                    "平衡准确率", "Cohen κ", "耗时 (ms)", "划分签名"]
    folds.append(fold_headers)
    style_header(folds, 1, len(fold_headers))
    for row in rows:
        for split in row.get("splits") or []:
            info_row = split.get("split") or {}
            folds.append([
                row.get("pipeline"),
                row.get("pipeline_name") or row.get("pipeline"),
                info_row.get("fold"),
                info_row.get("repeat"),
                info_row.get("seed"),
                info_row.get("n_train"),
                info_row.get("n_test"),
                round(float(split.get("balanced_accuracy") or 0.0), 6),
                round(float(split.get("cohen_kappa") or 0.0), 6),
                round(float(split.get("duration_ms") or 0.0), 2),
                _safe(str(info_row.get("signature", ""))),
            ])
    for row in folds.iter_rows(min_row=2, max_row=folds.max_row, max_col=len(fold_headers)):
        for cell in row:
            cell.border = border
        row[7].number_format = "0.0000"
        row[8].number_format = "0.0000"
    autosize(folds, {1: 18, 2: 28, 3: 8, 4: 9, 5: 8, 6: 10, 7: 10, 8: 14, 9: 12, 10: 12, 11: 42})
    folds.freeze_panes = "A2"
    if folds.max_row > 1:
        folds.conditional_formatting.add(
            f"H2:H{folds.max_row}",
            ColorScaleRule(start_type="num", start_value=0.5, start_color="F8D7DA",
                           end_type="num", end_value=1.0, end_color="C6E7D4"),
        )

    tests = book.create_sheet("统计检验")
    tests["A1"] = "统计检验"
    tests["A1"].font = title_font
    friedman = stats.get("friedman") or {}
    for column, value in enumerate(["Friedman", "统计量", "p 值", "显著"], start=1):
        tests.cell(row=3, column=column, value=value)
    style_header(tests, 3, 4)
    for column, value in enumerate([
        "跨算法差异",
        friedman.get("statistic"),
        friedman.get("p_value"),
        "是" if friedman.get("significant") else "否" if friedman.get("available") else "不可用",
    ], start=1):
        cell = tests.cell(row=4, column=column, value=value)
        cell.border = border
    tests["B4"].number_format = "0.0000"
    tests["C4"].number_format = "0.00E+00"
    row_cursor = 6
    nemenyi = stats.get("nemenyi") or {}
    if nemenyi.get("available"):
        tests.cell(row=row_cursor, column=1, value="Nemenyi 临界距离 CD")
        tests.cell(row=row_cursor, column=2, value=nemenyi.get("critical_difference")).number_format = "0.0000"
        row_cursor += 2
    pairwise = stats.get("pairwise") or []
    if pairwise:
        headers = ["算法 A", "算法 B", "统计量", "原始 p", "Holm 校正 p", "显著", "平均 BA 差"]
        for index, value in enumerate(headers, start=1):
            tests.cell(row=row_cursor, column=index, value=value)
        style_header(tests, row_cursor, len(headers))
        for pair in pairwise:
            row_cursor += 1
            values = [
                pair.get("pipeline_a"), pair.get("pipeline_b"), pair.get("statistic"),
                pair.get("p_value"), pair.get("p_value_adjusted"),
                "是" if pair.get("significant") else "否", pair.get("mean_diff"),
            ]
            for index, value in enumerate(values, start=1):
                cell = tests.cell(row=row_cursor, column=index, value=value)
                cell.border = border
                if index in (4, 5):
                    cell.number_format = "0.00E+00"
                if index in (3, 7):
                    cell.number_format = "0.0000"
    autosize(tests, {1: 22, 2: 22, 3: 14, 4: 14, 5: 16, 6: 8, 7: 14})

    calibration_sheet = book.create_sheet("概率校准")
    calibration_headers = ["算法 ID", "算法名称", "ECE", "MCE", "Brier", "NLL",
                           "准确率", "平均置信度", "评估试次"]
    calibration_sheet.append(calibration_headers)
    style_header(calibration_sheet, 1, len(calibration_headers))
    for row in rows:
        calibration = row.get("calibration") or {}
        if not calibration.get("available"):
            continue
        calibration_sheet.append([
            row.get("pipeline"),
            row.get("pipeline_name") or row.get("pipeline"),
            calibration.get("ece"),
            calibration.get("mce"),
            calibration.get("brier"),
            calibration.get("nll"),
            calibration.get("accuracy"),
            calibration.get("confidence_mean"),
            calibration.get("n_evaluated"),
        ])
    for row in calibration_sheet.iter_rows(
        min_row=2, max_row=calibration_sheet.max_row, max_col=len(calibration_headers)
    ):
        for cell in row:
            cell.border = border
        for index in (3, 4, 5, 6, 7, 8):
            row[index - 1].number_format = "0.0000"
    autosize(calibration_sheet, {1: 20, 2: 30, 3: 10, 4: 10, 5: 10, 6: 10, 7: 10, 8: 12, 9: 10})
    calibration_sheet.freeze_panes = "A2"

    reliability_sheet = book.create_sheet("可靠性曲线")
    reliability_headers = ["算法 ID", "算法名称", "区间下限", "区间上限",
                           "平均置信度", "实际正确率", "试次数"]
    reliability_sheet.append(reliability_headers)
    style_header(reliability_sheet, 1, len(reliability_headers))
    for row in rows:
        calibration = row.get("calibration") or {}
        for item in calibration.get("reliability") or []:
            reliability_sheet.append([
                row.get("pipeline"),
                row.get("pipeline_name") or row.get("pipeline"),
                item.get("bin_lower"),
                item.get("bin_upper"),
                item.get("confidence"),
                item.get("accuracy"),
                item.get("count"),
            ])
    for row in reliability_sheet.iter_rows(
        min_row=2, max_row=reliability_sheet.max_row, max_col=len(reliability_headers)
    ):
        for cell in row:
            cell.border = border
        for index in (5, 6):
            row[index - 1].number_format = "0.0000"
    autosize(reliability_sheet, {1: 20, 2: 30, 3: 10, 4: 10, 5: 12, 6: 12, 7: 9})
    reliability_sheet.freeze_panes = "A2"

    label_names = manifest.get("label_names") or []
    matrix_sheet = book.create_sheet("混淆矩阵")
    row_cursor = 1
    for row in rows:
        splits = row.get("splits") or []
        y_true = [label for split in splits for label in split.get("y_true", [])]
        y_pred = [label for split in splits for label in split.get("y_pred", [])]
        if not y_true or len(y_true) != len(y_pred):
            continue
        classes = sorted(set(y_true) | set(y_pred))
        matrix = confusion_matrix(np.asarray(y_true), np.asarray(y_pred), len(classes))
        labels = [
            label_names[index] if index < len(label_names) else f"类别 {index}"
            for index in classes
        ]
        matrix_sheet.cell(row=row_cursor, column=1, value=row.get("pipeline_name") or row.get("pipeline")).font = Font(bold=True, color="1F3A5F")
        row_cursor += 1
        matrix_sheet.cell(row=row_cursor, column=1, value="真实 \\ 预测")
        for index, label in enumerate(labels, start=2):
            matrix_sheet.cell(row=row_cursor, column=index, value=_safe(str(label)))
        style_header(matrix_sheet, row_cursor, len(labels) + 1)
        for matrix_row, label in enumerate(labels):
            row_cursor += 1
            matrix_sheet.cell(row=row_cursor, column=1, value=_safe(str(label)))
            for matrix_column in range(len(labels)):
                cell = matrix_sheet.cell(
                    row=row_cursor, column=matrix_column + 2,
                    value=int(matrix[matrix_row, matrix_column]),
                )
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
        row_cursor += 2
    autosize(matrix_sheet, {1: 24, 2: 14, 3: 14, 4: 14, 5: 14, 6: 14, 7: 14, 8: 14, 9: 14})

    xai_items = _explanation_items(detail)
    xai_sheet = book.create_sheet("模型在看哪里")
    xai_headers = ["解释 ID", "算法", "Fold", "Repeat", "Seed", "平衡准确率",
                   "删除 AUC", "插入 AUC", "随机删除 AUC", "AOPC", "closure",
                   "半折稳定性 ρ", "归一化可用", "警告"]
    xai_sheet.append(xai_headers)
    style_header(xai_sheet, 1, len(xai_headers))
    for item in xai_items:
        document = item.get("document") or {}
        summary = item.get("summary") or {}
        faith = document.get("faithfulness") or {}
        audit = document.get("audit") or {}
        xai_sheet.append([
            item.get("id"),
            item.get("pipeline_name") or item.get("pipeline"),
            item.get("fold"),
            item.get("repeat"),
            item.get("seed"),
            (summary.get("performance") or {}).get("balanced_accuracy"),
            faith.get("deletion_auc"),
            faith.get("insertion_auc"),
            faith.get("random_deletion_auc"),
            faith.get("aopc"),
            audit.get("closure_ratio"),
            (faith.get("stability") or {}).get("half_split_spearman"),
            "是" if (faith.get("normalization") or {}).get("usable", True) else "否",
            _safe("；".join(summary.get("warnings") or [])),
        ])
    for row in xai_sheet.iter_rows(
        min_row=2, max_row=xai_sheet.max_row, max_col=len(xai_headers)
    ):
        for cell in row:
            cell.border = border
        for index in (6, 7, 8, 9, 10, 11, 12):
            row[index - 1].number_format = "0.0000"
    autosize(xai_sheet, {1: 10, 2: 26, 3: 8, 4: 9, 5: 8, 6: 12, 7: 11, 8: 11,
                         9: 14, 10: 10, 11: 10, 12: 12, 13: 12, 14: 60})
    xai_sheet.freeze_panes = "A2"

    xai_channel_sheet = book.create_sheet("解释通道重要性")
    channel_headers = ["解释 ID", "算法", "通道", "遮挡 Δlog-odds",
                       "内在参数", "频带消融", "梯度幅值", "频带能量"]
    xai_channel_sheet.append(channel_headers)
    style_header(xai_channel_sheet, 1, len(channel_headers))
    for item in xai_items:
        for row in _xai_channel_rows(item, limit=10_000, formatted=False):
            xai_channel_sheet.append(
                [
                    item.get("id"),
                    item.get("pipeline_name") or item.get("pipeline"),
                    *row,
                ]
            )
    for row in xai_channel_sheet.iter_rows(
        min_row=2, max_row=xai_channel_sheet.max_row, max_col=len(channel_headers)
    ):
        for cell in row:
            cell.border = border
        for index in (4, 5, 6, 7, 8):
            if isinstance(row[index - 1].value, (int, float)):
                row[index - 1].number_format = "0.0000"
    autosize(xai_channel_sheet, {1: 10, 2: 26, 3: 12, 4: 16, 5: 12, 6: 12, 7: 12, 8: 12})
    xai_channel_sheet.freeze_panes = "A2"

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


__all__ = ["build_docx", "build_html_report", "build_latex", "build_workbook"]


REFERENCES: tuple[dict[str, str], ...] = (
    {
        "key": "neurokiln",
        "text": "NeuroKiln Contributors. NeuroKiln: 脑机接口算法评测与超参数搜索平台[CP/OL]. 2026.",
    },
    {
        "key": "moabb",
        "text": "JAYARAM V, BARACHANT A. MOABB: trustworthy algorithm benchmarking for BCIs[J]. "
                "Journal of Neural Engineering, 2018, 15(6): 066011.",
    },
    {
        "key": "mne",
        "text": "GRAMFORT A, LUESSI M, LARSON E, et al. MEG and EEG data analysis with "
                "MNE-Python[J]. Frontiers in Neuroscience, 2013, 7: 267.",
    },
    {
        "key": "scikit",
        "text": "PEDREGOSA F, VAROQUAUX G, GRAMFORT A, et al. Scikit-learn: machine learning "
                "in Python[J]. Journal of Machine Learning Research, 2011, 12: 2825-2830.",
    },
    {
        "key": "pyriemann",
        "text": "BARACHANT A, BONNET S, CONGEDO M, et al. Classification of covariance matrices "
                "using a Riemannian-based kernel for BCI applications[J]. Neurocomputing, 2013, "
                "112: 172-178.",
    },
    {
        "key": "demsar",
        "text": "DEMSAR J. Statistical comparisons of classifiers over multiple data sets[J]. "
                "Journal of Machine Learning Research, 2006, 7: 1-30.",
    },
)


def _reference_entries(detail: dict[str, Any]) -> list[str]:
    entries = [item["text"] for item in REFERENCES]
    dataset = detail.get("dataset") or {}
    manifest = dataset.get("manifest") or {}
    citation = str(manifest.get("citation") or "").strip()
    if citation and citation not in entries:
        entries.insert(1, citation)
    return entries


def _ranked_rows(detail: dict[str, Any]) -> list[dict[str, Any]]:
    return sorted(
        detail.get("rows") or [],
        key=lambda row: float(row.get("balanced_accuracy_mean") or 0.0),
        reverse=True,
    )


def _row_duration(row: dict[str, Any]) -> float:
    durations = [float(item.get("duration_ms", 0.0)) for item in row.get("splits") or []]
    return sum(durations) / len(durations) if durations else 0.0


def _statistics_block(detail: dict[str, Any]) -> dict[str, Any]:
    return detail.get("statistics") or (detail.get("summary") or {}).get("statistics") or {}


def _document_meta(detail: dict[str, Any]) -> list[tuple[str, str]]:
    dataset = detail.get("dataset") or {}
    manifest = dataset.get("manifest") or {}
    return [
        ("实验", f"#{detail.get('id')} · {detail.get('name') or ''}"),
        ("数据集", f"{dataset.get('name') or dataset.get('id') or '—'} · v{dataset.get('version') or '—'}"),
        ("被试", ", ".join(str(item) for item in (manifest.get("subjects") or [])) or "—"),
        ("试次 / 通道 / 采样率", f"{manifest.get('loaded_n_trials') or '—'} / "
                              f"{manifest.get('loaded_channels') or '—'} / "
                              f"{manifest.get('loaded_sfreq') or '—'} Hz"),
        ("协议", str(detail.get("protocol") or "—")),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]


def _docx_style(document) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        try:
            style = document.styles[name]
        except KeyError:
            continue
        style.font.name = "Microsoft YaHei"
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            from docx.oxml import OxmlElement

            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.styles["Normal"].font.size = Pt(10.5)


def _docx_table(document, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> Any:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, label in enumerate(headers):
        cell = header_cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(str(label))
        run.bold = True
        run.font.size = Pt(9)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "EAF1F8")
        cell._tc.get_or_add_tcPr().append(shading)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run("" if value is None else str(value))
            run.font.size = Pt(9)
    document.add_paragraph()
    return table


def build_docx(detail: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    _docx_style(document)
    rows = _ranked_rows(detail)
    stats = _statistics_block(detail)

    title = document.add_heading(f"NeuroKiln 实验报告 · #{detail.get('id')}", level=0)
    for run in title.runs:
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(22)
    subtitle = document.add_paragraph(detail.get("name") or "未命名实验")
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].bold = True

    document.add_heading("实验概览", level=1)
    _docx_table(document, ["项目", "内容"], _document_meta(detail))

    if rows:
        best = rows[0]
        friedman = stats.get("friedman") or {}
        overview = [
            ["最佳算法", best.get("pipeline_name") or best.get("pipeline"),
             f"BA {float(best.get('balanced_accuracy_mean') or 0.0):.4f}"],
            ["Friedman p", _p_value(friedman.get("p_value")),
             "差异显著" if friedman.get("significant") else "未达显著"],
            ["统计单位", "被试" if stats.get("unit") == "subject" else "折",
             f"{stats.get('n_units', 0)} 个单位 · {stats.get('n_splits', 0)} 折"],
            ["总耗时", f"{float(detail.get('duration_ms') or 0.0) / 1000.0:.1f} s",
             f"引擎 {detail.get('engine_version', '—')}"],
        ]
        _docx_table(document, ["指标", "数值", "说明"], overview)

    document.add_heading("算法排名", level=1)
    ranking = []
    for rank, row in enumerate(rows, start=1):
        ranking.append([
            rank,
            row.get("pipeline_name") or row.get("pipeline"),
            f"{float(row.get('balanced_accuracy_mean') or 0.0):.4f}",
            f"± {float(row.get('balanced_accuracy_std') or 0.0):.4f}",
            _ci(row.get("balanced_accuracy_ci95")),
            f"{float(row.get('cohen_kappa_mean') or 0.0):.4f}",
            _num((row.get("calibration") or {}).get("ece")),
            f"{_row_duration(row):.1f} ms",
            len(row.get("splits") or []),
        ])
    _docx_table(
        document,
        ["#", "算法", "平衡准确率", "标准差", "95% CI", "Cohen κ", "ECE", "单折耗时", "折数"],
        ranking,
    )

    document.add_heading("统计检验", level=1)
    friedman = stats.get("friedman") or {}
    if friedman.get("available"):
        document.add_paragraph(
            f"Friedman χ² = {_num(friedman.get('statistic'), 3)}，"
            f"p = {_p_value(friedman.get('p_value'))}。"
        )
    else:
        document.add_paragraph(str(friedman.get("reason") or "该实验未生成 Friedman 检验。"))
    nemenyi = stats.get("nemenyi") or {}
    if nemenyi.get("available"):
        document.add_paragraph(
            f"Nemenyi 临界距离 CD = {_num(nemenyi.get('critical_difference'), 3)}"
            f"（α={stats.get('alpha', 0.05)}）。"
        )
    pairwise = stats.get("pairwise") or []
    if pairwise:
        _docx_table(
            document,
            ["算法对", "统计量", "原始 p", "Holm 校正 p", "显著"],
            [
                [
                    f"{item.get('pipeline_a')} vs {item.get('pipeline_b')}",
                    _num(item.get("statistic"), 3),
                    _p_value(item.get("p_value")),
                    _p_value(item.get("p_value_adjusted")),
                    "是" if item.get("significant") else "否",
                ]
                for item in pairwise
            ],
        )

    label_names = ((detail.get("dataset") or {}).get("manifest") or {}).get("label_names") or []
    matrix_rows: list[list[Any]] = []
    for row in rows:
        splits = row.get("splits") or []
        y_true = [label for split in splits for label in split.get("y_true", [])]
        y_pred = [label for split in splits for label in split.get("y_pred", [])]
        if not y_true or len(y_true) != len(y_pred):
            continue
        classes = sorted(set(y_true) | set(y_pred))
        matrix = confusion_matrix(np.asarray(y_true), np.asarray(y_pred), len(classes))
        matrix_rows.append([
            row.get("pipeline_name") or row.get("pipeline"),
            " / ".join(
                label_names[index] if index < len(label_names) else f"类别 {index}"
                for index in classes
            ),
            "; ".join(
                "[" + ", ".join(str(int(value)) for value in matrix_line) + "]"
                for matrix_line in matrix
            ),
        ])
    if matrix_rows:
        document.add_heading("混淆矩阵（全折汇总）", level=1)
        _docx_table(document, ["算法", "类别顺序", "矩阵行（真实 × 预测）"], matrix_rows)

    xai_items = _explanation_items(detail)
    if xai_items:
        from docx.shared import Cm

        document.add_heading("模型在看哪里（可解释性）", level=1)
        document.add_paragraph(
            "遮挡基线取训练折逐通道均值，测试折不参与任何拟合；"
            "归因目标是模型自身预测类别的 log-odds。"
        )
        for item in xai_items:
            document.add_heading(_xai_label(item), level=2)
            _docx_table(
                document,
                ["指标", "数值", "说明"],
                [[label, value, hint] for label, value, hint in _xai_metrics(item)],
            )
            document_data = item.get("document") or {}
            images = item.get("images") or {}
            figure_specs = [
                ((document_data.get("occlusion") or {}).get("topomap"), "遮挡归因（Δlog-odds）"),
                ((document_data.get("bands") or {}).get("topomap"), "频带消融"),
                ((document_data.get("gradient") or {}).get("topomap"), "梯度归因"),
            ]
            for asset, caption in figure_specs:
                uri = images.get(asset) if asset else None
                if not uri or not uri.startswith("data:image/png;base64,"):
                    continue
                try:
                    payload = base64.b64decode(uri.split(",", 1)[1])
                except (ValueError, IndexError):
                    continue
                document.add_picture(BytesIO(payload), width=Cm(5.4))
                note = document.add_paragraph(caption)
                for run in note.runs:
                    run.font.size = Pt(9)
            channel_rows = _xai_channel_rows(item)
            if channel_rows:
                _docx_table(
                    document,
                    ["通道", "遮挡 Δ", "内在参数", "频带消融", "梯度", "频带能量"],
                    channel_rows,
                )
            for message in (item.get("summary") or {}).get("warnings") or []:
                paragraph = document.add_paragraph(f"注意：{message}")
                for run in paragraph.runs:
                    run.font.size = Pt(9)

    document.add_heading("参考文献", level=1)
    for index, entry in enumerate(_reference_entries(detail), start=1):
        paragraph = document.add_paragraph(f"[{index}] {entry}")
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.size = Pt(9)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_LATEX_ESCAPES = {
    "&": r"\&",
    "%": r"\%",
    "$": r"\$",
    "#": r"\#",
    "_": r"\_",
    "{": r"\{",
    "}": r"\}",
}


def _tex(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\\", "\x00")
    for source, target in _LATEX_ESCAPES.items():
        text = text.replace(source, target)
    text = text.replace("~", r"\textasciitilde{}").replace("^", r"\textasciicircum{}")
    return text.replace("\x00", r"\textbackslash{}")


def _tex_cell(value: Any) -> str:
    return _tex(value).replace("\n", " ")


def _tex_longtable(headers: Sequence[str], rows: Sequence[Sequence[Any]], spec: str) -> str:
    head = " & ".join(rf"\textbf{{{_tex_cell(item)}}}" for item in headers) + r" \\"
    body = "\n".join(
        " & ".join(_tex_cell(value) for value in row) + r" \\" for row in rows
    )
    return (
        rf"\begin{{longtable}}{{{spec}}}" "\n"
        r"\toprule" "\n" + head + "\n" + r"\midrule" "\n" + r"\endhead" "\n"
        + body + "\n" + r"\bottomrule" "\n" + r"\end{longtable}" "\n"
    )


def build_latex(detail: dict[str, Any]) -> str:
    rows = _ranked_rows(detail)
    stats = _statistics_block(detail)
    dataset = detail.get("dataset") or {}
    manifest = dataset.get("manifest") or {}
    title = f"NeuroKiln 实验报告 · #{detail.get('id')}"
    generated = datetime.now().strftime("%Y-%m-%d %H:%M")

    ranking = [
        [
            rank,
            row.get("pipeline_name") or row.get("pipeline"),
            f"{float(row.get('balanced_accuracy_mean') or 0.0):.4f}",
            f"{float(row.get('balanced_accuracy_std') or 0.0):.4f}",
            _ci(row.get("balanced_accuracy_ci95")),
            f"{float(row.get('cohen_kappa_mean') or 0.0):.4f}",
            f"{_row_duration(row):.1f}",
        ]
        for rank, row in enumerate(rows, start=1)
    ]
    friedman = stats.get("friedman") or {}
    nemenyi = stats.get("nemenyi") or {}
    pairwise = stats.get("pairwise") or []
    pairwise_rows = [
        [
            f"{item.get('pipeline_a')} vs {item.get('pipeline_b')}",
            _num(item.get("statistic"), 3),
            _p_value(item.get("p_value")),
            _p_value(item.get("p_value_adjusted")),
            "是" if item.get("significant") else "否",
        ]
        for item in pairwise
    ]

    parts = [
        "% 使用 XeLaTeX 编译：xelatex report.tex",
        r"\documentclass[UTF8,a4paper,11pt]{ctexart}",
        r"\usepackage[margin=2.2cm]{geometry}",
        r"\usepackage{booktabs}",
        r"\usepackage{longtable}",
        r"\usepackage{array}",
        r"\usepackage{hyperref}",
        rf"\title{{{_tex(title)}}}",
        r"\author{NeuroKiln 脑炉}",
        rf"\date{{{generated}}}",
        r"\begin{document}",
        r"\maketitle",
        r"\section*{实验概览}",
        r"\begin{tabular}{@{}ll@{}}",
        r"\toprule",
    ]
    for label, value in _document_meta(detail):
        parts.append(rf"\textbf{{{_tex_cell(label)}}} & {_tex_cell(value)} \\")
    parts += [
        r"\bottomrule",
        r"\end{tabular}",
        "",
        r"\noindent 数据签名：\texttt{" + _tex_cell(dataset.get("signature") or "—") + "}。",
        "",
        r"\section*{算法排名}",
        _tex_longtable(
            ["#", "算法", "平衡准确率", "标准差", "95% CI", "Cohen κ", "单折耗时 (ms)"],
            ranking,
            r"rlrrrrr",
        ),
        r"\section*{统计检验}",
    ]
    if friedman.get("available"):
        parts.append(
            r"\noindent Friedman $\chi^2 = "
            + _num(friedman.get("statistic"), 3)
            + "$，$p = "
            + _p_value(friedman.get("p_value")).replace("%", r"\%")
            + "$。"
        )
    else:
        parts.append(r"\noindent " + _tex_cell(friedman.get("reason") or "该实验未生成 Friedman 检验。"))
    if nemenyi.get("available"):
        parts.append(
            "\n" + r"\noindent Nemenyi 临界距离 CD = "
            + _num(nemenyi.get("critical_difference"), 3)
            + rf"，$\alpha = {_num(stats.get('alpha', 0.05), 2)}$。"
        )
    if pairwise_rows:
        parts.append(_tex_longtable(
            ["算法对", "统计量", "原始 $p$", "Holm 校正 $p$", "显著"],
            pairwise_rows,
            r"lrrrr",
        ))
    xai_items = _explanation_items(detail)
    if xai_items:
        xai_rows = []
        for item in xai_items:
            document_data = item.get("document") or {}
            faith = document_data.get("faithfulness") or {}
            audit = document_data.get("audit") or {}
            xai_rows.append([
                item.get("pipeline_name") or item.get("pipeline"),
                item.get("fold", 0),
                _num(faith.get("deletion_auc")),
                _num(faith.get("insertion_auc")),
                _num(faith.get("random_deletion_auc")),
                _num(faith.get("aopc")),
                _num(audit.get("closure_ratio")),
            ])
        parts += [
            r"\section*{模型在看哪里（可解释性）}",
            r"\noindent 遮挡基线取训练折逐通道均值，测试折不参与任何拟合；"
            r"归因目标是模型自身预测类别的 log-odds。删除曲线 AUC 越低、插入曲线 AUC 越高越忠实。"
            r"地形图见 HTML/Web 报告。",
            _tex_longtable(
                ["算法", "Fold", "删除 AUC", "插入 AUC", "随机删除 AUC", "AOPC", "closure"],
                xai_rows,
                r"lcrrrrr",
            ),
        ]
    parts += [r"\section*{参考文献}", r"\begin{thebibliography}{99}"]
    for item in _reference_entries(detail):
        parts.append(r"\bibitem{} " + _tex_cell(item))
    parts += [r"\end{thebibliography}", r"\end{document}", ""]
    return "\n".join(parts)

